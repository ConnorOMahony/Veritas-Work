"""
Veritas Work - Master Investment Summary & Cash Flow Builder

For each client, this script:
  1. Reads all matched investment statements - both .xlsx AND .pdf - from
     the folder (transaction sheets).
  2. Finds that client's EXISTING Investment Summary file already sitting
     in the folder. If none exists, defaults to the blank template.
  3. Opens that existing file exactly as it is.
  4. Duplicates the pristine "DDMMYY" template tab to ensure perfect layouts.
  5. Updates the account values on ONLY that new tab. 
  6. Updates the existing "Cash Flows" tab IN PLACE to preserve images/logos, 
     wiping old data and rebuilding it cleanly.
  7. Saves the result as a new file. The original files on disk are never modified.
"""

import os
import re
import io
import copy
import zipfile
import logging
from datetime import date, datetime
from collections import defaultdict

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# CONFIG - MAIN DASHBOARD
# ---------------------------------------------------------------------------
logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')

try:
    FOLDER = os.path.dirname(os.path.abspath(__file__))
except NameError:
    FOLDER = os.getcwd()

OUTPUT_PREFIX = "Investment_Summary_"
TEMPLATE_SHEET_NAME = "DDMMYY"
CLOSING_BALANCE_LABEL = "Closing Balance"

CLIENT_NAME_COLUMN = 1   # column A
INVESTMENTS_COLUMN = 2   # column B
INCEPTION_DATE_COLUMN = 3  # column C
NET_CONTRIB_COLUMN = 4   # column D
CURRENT_VALUE_COLUMN = 5  # column E
TITLE_ROW = 2
TITLE_COLUMN = 1

COL_REPORT_DATE = 0
COL_CLIENT_NAME = 4
COL_CONTRACT_NUMBER = 5
COL_VEHICLE_NAME = 7
COL_HOLDING = 9
COL_TRANSACTION_DATE = 10
COL_TRANSACTION_TYPE = 11
COL_TRANSACTION_AMOUNT = 14

PLACEHOLDER_ROWS = {
    5:  ("Living Annuity", ["living annuity"], []),
    7:  ("Retirement Annuity", ["retirement annuity"], ["living", "provident", "preservation"]),
    9:  ("Provident/Preservation Fund", ["provident", "preservation"], []),
    13: ("Investment Platform Unit Trust", ["unit trust", "investment platform"], []),
    16: ("Tax-Free Investment", ["tax-free", "tax free"], []),
    21: ("Offshore Investment", ["offshore"], ["momentum", "international"]),
    23: ("Momentum Wealth International", ["momentum", "international"], []),
}

PLACEHOLDER_SPANS = {
    5: 2, 7: 2, 9: 2, 13: 3, 16: 2, 21: 2, 23: 2
}

COLUMN_HEADER_ALIASES = {
    "report_date": ["report date", "statement date", "valuation date"],
    "client_name": ["client name", "investor name", "member name"],
    "contract_number": ["investment contract number", "contract number",
                        "policy number", "account number"],
    "vehicle_name": ["investment vehicle name", "product name",
                     "investment name", "product"],
    "holding": ["holdings", "holding", "fund name", "fund"],
    "transaction_date": ["transaction date"],
    "transaction_type": ["transaction type"],
    "transaction_amount": ["transaction amount", "amount"],
}

FALLBACK_POSITIONS = {
    "report_date": COL_REPORT_DATE,
    "client_name": COL_CLIENT_NAME,
    "contract_number": COL_CONTRACT_NUMBER,
    "vehicle_name": COL_VEHICLE_NAME,
    "holding": COL_HOLDING,
    "transaction_date": COL_TRANSACTION_DATE,
    "transaction_type": COL_TRANSACTION_TYPE,
    "transaction_amount": COL_TRANSACTION_AMOUNT,
}

OPTIONAL_COLUMN_ALIASES = {
    "inception_date": ["inception date", "investment start date", "start date",
                        "commencement date", "policy start date"],
}

CONTRIBUTION_KEYWORDS = ["contribution", "deposit", "debit order", "premium", "recurring investment"]
WITHDRAWAL_KEYWORDS = ["withdrawal", "disinvestment", "surrender", "once-off withdrawal"]

DATE_STRING_FORMATS = [
    "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y",
    "%d %B %Y", "%d %b %Y", "%m/%d/%Y",
]

# ---------------------------------------------------------------------------
# CONFIG - CASH FLOW ENGINE
# ---------------------------------------------------------------------------
CF_DESCRIPTION_MAPPING = {
    "investment": "Additional investment",
    "recurring lumpsum contribution": "Additional investment",
    "switch in": "Additional investment (trf in)",
    "switch out": "Withdrawal (trf out)",
    "withdrawal": "Partial surrender",
    "partial surrender": "Partial surrender"
}

CF_IGNORE_KEYWORDS = [
    "model rebalance", "fee", "interest", "reinvestment", "discretionary",
    "opening market value", "market value", "total value", "transactions in progress",
    "savings withdrawable", "distributions earned"
]

CF_ACCOUNT_MAPPING = {
    "AGRA": {"product": "Retirement Annuity", "category": "Retirement Funds"},
    "AGLA": {"product": "Living Annuity", "category": "Retirement Funds"},
    "AGTF": {"product": "Tax-Free Investment", "category": "Discretionary Investments"},
    "AGLP": {"product": "Investment Platform Unit Trust", "category": "Discretionary Investments"},
    "AGEN": {"product": "Endowment", "category": "Discretionary Investments"},
    "AGUP": {"product": "Umbrella Provident Fund", "category": "Retirement Funds"},
    "AGPP": {"product": "Preservation Provident Fund", "category": "Retirement Funds"},
    "AGEP": {"product": "Preservation Pension Fund", "category": "Retirement Funds"},
    "AGOP": {"product": "Offshore Investment Platform", "category": "Offshore Investments"}
}

CF_CATEGORY_ORDER = ["Retirement Funds", "Discretionary Investments", "Offshore Investments"]


# ---------------------------------------------------------------------------
# SHARED UTILITIES
# ---------------------------------------------------------------------------

def extract_zip_files():
    extracted_any = False
    for root, dirs, files in os.walk(FOLDER):
        for name in files:
            if not name.lower().endswith(".zip"):
                continue
            zip_path = os.path.join(root, name)
            try:
                with zipfile.ZipFile(zip_path) as zf:
                    zf.extractall(FOLDER)
                logging.info(f"Extracted zip file: {name}")
                extracted_any = True
            except Exception as exc:
                logging.warning(f"Could not extract '{name}': {exc}")
    return extracted_any

def normalize(name):
    return re.sub(r"[\s_\-]+", "", name.lower())

def normalize_client_name(name):
    if not name: return "Unknown Client"
    name = str(name).strip().upper()
    name = re.sub(r"^(MR|MRS|MS|DR|PROF)\.?\s+", "", name)
    name = re.sub(r"\s+", " ", name)
    return name

def parse_date_value(value):
    if value is None: return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    text = str(value).strip()
    if not text: return None
    for fmt in DATE_STRING_FORMATS:
        try: return datetime.strptime(text, fmt).date()
        except ValueError: continue
    return None

def is_investment_summary_filename(name):
    if not name.lower().endswith(".xlsx") or name.startswith("~$"):
        return False
    return "investmentsummary" in normalize(name)

def find_investment_summary_files():
    found = []
    for root, dirs, files in os.walk(FOLDER):
        for name in files:
            if is_investment_summary_filename(name):
                found.append(os.path.join(root, name))
    return found

def find_statement_files(summary_files):
    summary_abspaths = {os.path.abspath(f) for f in summary_files}
    files = []
    for root, dirs, files_list in os.walk(FOLDER):
        for name in files_list:
            lower = name.lower()
            if not (lower.endswith(".xlsx") or lower.endswith(".pdf")): continue
            if name.startswith("~$"): continue
            if is_investment_summary_filename(name): continue
            full_path = os.path.join(root, name)
            if os.path.abspath(full_path) in summary_abspaths: continue
            files.append(full_path)
    return files

def match_summary_file_for_client(client_name, summary_files):
    specific_files = [f for f in summary_files if "template" not in normalize(os.path.basename(f))]
    template_files = [f for f in summary_files if "template" in normalize(os.path.basename(f))]

    non_dated_specific = [f for f in specific_files if not re.search(r"_\d{6}\.xlsx$", f, re.IGNORECASE)]
    candidate_pool = non_dated_specific if non_dated_specific else specific_files

    if len(candidate_pool) == 1 and not template_files:
        return candidate_pool[0]

    tokens = [t for t in re.split(r"\s+", client_name) if len(t) > 1]
    best_file, best_score = None, 0
    for f in candidate_pool:
        norm_f = normalize(os.path.basename(f))
        score = sum(1 for t in tokens if normalize(t) in norm_f)
        if score > best_score:
            best_score, best_file = score, f

    if best_file is not None:
        return best_file

    if template_files:
        return template_files[0]
        
    if summary_files:
        return summary_files[0]

    return None

def parse_money(value):
    if value is None: return 0.0
    if isinstance(value, (int, float)): return float(value)
    text = str(value).strip().replace('\xa0', ' ').replace("R", "").replace(",", "").replace(" ", "")
    if text in ("", "-"): return 0.0
    if text.startswith("(") and text.endswith(")"): text = "-" + text[1:-1]
    try: return float(text)
    except ValueError: return 0.0


# ---------------------------------------------------------------------------
# DASHBOARD: DATA EXTRACTION ENGINE 
# ---------------------------------------------------------------------------

def find_header_row(ws, max_scan=20):
    all_aliases = [a for aliases in COLUMN_HEADER_ALIASES.values() for a in aliases]
    best_row_idx, best_score = 1, -1
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        if row is None: continue
        headers = [str(h).strip().lower() if h is not None else "" for h in row]
        score = sum(1 for alias in all_aliases if any(alias in h for h in headers))
        if score > best_score:
            best_score, best_row_idx = score, idx
    return best_row_idx, best_score

def build_column_map(ws):
    header_row_idx, header_score = find_header_row(ws)
    header_row = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True), [])
    headers = [str(h).strip().lower() if h is not None else "" for h in header_row]

    column_map = {}
    warnings = [f"Detected header row at row {header_row_idx} (match score {header_score}): {list(header_row)}"]
    for role, aliases in COLUMN_HEADER_ALIASES.items():
        found_index = None
        for alias in aliases:
            for i, h in enumerate(headers):
                if alias in h:
                    found_index = i
                    break
            if found_index is not None: break
        if found_index is None:
            found_index = FALLBACK_POSITIONS[role]
            warnings.append(f"Header '{role}' not found - falling back to col {found_index + 1}.")
        column_map[role] = found_index

    for role, aliases in OPTIONAL_COLUMN_ALIASES.items():
        found_index = None
        for alias in aliases:
            for i, h in enumerate(headers):
                if alias in h:
                    found_index = i
                    break
            if found_index is not None: break
        column_map[role] = found_index
        if found_index is None:
            warnings.append(f"Optional header '{role}' not found - will be left blank/unchanged.")

    return column_map, warnings

def is_allan_gray_fund_summary(ws):
    try:
        for row in ws.iter_rows(min_row=1, max_row=20, max_col=6, values_only=True):
            for val in row:
                if val is not None:
                    text = str(val).strip().lower()
                    if "investor" in text or "allan gray" in text:
                        return True
    except Exception: pass
    return False

def read_old_mutual_style_statement(ws, filepath):
    column_map, warnings = build_column_map(ws)
    max_col = max(v for v in column_map.values() if v is not None)
    latest = {}
    net_contributions = {}   
    inception_dates = {}     
    seen_txn_types = set()   

    inception_col = column_map.get("inception_date")

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or len(row) <= max_col: continue

        contract = row[column_map["contract_number"]]
        holding = row[column_map["holding"]]
        if not contract: continue
        key = (contract, holding)

        if inception_col is not None and inception_col < len(row):
            parsed_inc = parse_date_value(row[inception_col])
            if parsed_inc:
                existing_inc = inception_dates.get(key)
                if existing_inc is None or parsed_inc < existing_inc:
                    inception_dates[key] = parsed_inc

        txn_type_raw = row[column_map["transaction_type"]]
        if txn_type_raw is not None:
            seen_txn_types.add(str(txn_type_raw).strip())
        txn_type_norm = str(txn_type_raw).strip().lower() if txn_type_raw else ""
        amount_val = row[column_map["transaction_amount"]]
        if any(lbl in txn_type_norm for lbl in CONTRIBUTION_KEYWORDS):
            net_contributions[key] = net_contributions.get(key, 0.0) + abs(parse_money(amount_val))
        elif any(lbl in txn_type_norm for lbl in WITHDRAWAL_KEYWORDS):
            net_contributions[key] = net_contributions.get(key, 0.0) - abs(parse_money(amount_val))

        if txn_type_raw != CLOSING_BALANCE_LABEL: continue

        t_date = row[column_map["transaction_date"]]
        amount = amount_val or 0
        client_name = row[column_map["client_name"]]
        vehicle_name = row[column_map["vehicle_name"]]
        report_date = row[column_map["report_date"]]

        existing = latest.get(key)
        if existing is None or (t_date and existing[0] and t_date > existing[0]):
            latest[key] = (t_date, amount, client_name, vehicle_name, report_date, row)

    results = {}
    for (contract, holding), (t_date, amount, client_name, vehicle_name, report_date, row) in latest.items():
        contract_str = str(contract).strip()
        key = (contract, holding)
        if contract_str not in results:
            results[contract_str] = {
                "client_name": (client_name or "Unknown Client").strip(),
                "vehicle_name": vehicle_name,
                "value": 0.0,
                "report_date": report_date,
                "transaction_date": t_date,
                "net_contribution": 0.0,
                "inception_date": None,
                "funds": {}
            }
        
        val = parse_money(amount)
        results[contract_str]["value"] += val
        
        if holding:
            results[contract_str]["funds"][str(holding)] = val

        existing_t_date = results[contract_str].get("transaction_date")
        parsed_existing = parse_date_value(existing_t_date)
        parsed_new = parse_date_value(t_date)
        if parsed_new and (not parsed_existing or parsed_new > parsed_existing):
            results[contract_str]["transaction_date"] = t_date

        results[contract_str]["net_contribution"] += net_contributions.get(key, 0.0)

        inc_date = inception_dates.get(key)
        existing_inc = results[contract_str]["inception_date"]
        if inc_date and (existing_inc is None or inc_date < existing_inc):
            results[contract_str]["inception_date"] = inc_date

    fname = os.path.basename(filepath)
    for w in warnings:
        logging.info(f"{fname}: {w}")
    if seen_txn_types:
        logging.info(f"{fname}: Transaction Type values found: {sorted(seen_txn_types)}")
    if not any(net_contributions.values()):
        logging.info(f"{fname}: No 'Contribution'/'Withdrawal' rows matched - Net Contribution will show as 0.")
    if inception_col is None:
        logging.info(f"{fname}: No 'Inception Date' - left blank.")
    return results

def read_allan_gray_fund_summary(ws, filepath):
    client_name, vehicle_name, contract_number, report_date = None, None, None, None
    inception_date_value = None
    value_col, header_row_idx, initial_contrib_col = None, None, None
    total_value, total_value_including_progress = None, None
    total_initial_contribution = None
    funds = {}

    INVESTOR_LABELS = {"investor", "investor name", "client name", "member name"}
    PRODUCT_LABELS = {"product"}
    ACCOUNT_NUMBER_LABELS = {"account number", "policy number", "contract number"}
    INCEPTION_DATE_LABELS = {"inception date", "commencement date", "start date"}
    DATE_IN_TEXT = re.compile(r"\d{1,2}\s+[A-Za-z]+\s+\d{4}")

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    totals_row_idx = None

    for idx, row in enumerate(rows, start=1):
        if row is None: continue
        for c, val in enumerate(row):
            if val is None: continue
            text = str(val).strip().lower().rstrip(":").strip()
            if text in INVESTOR_LABELS and c + 1 < len(row) and row[c + 1] is not None:
                client_name = row[c + 1]
            elif text in PRODUCT_LABELS and c + 1 < len(row) and row[c + 1] is not None:
                vehicle_name = row[c + 1]
            elif text in ACCOUNT_NUMBER_LABELS and c + 1 < len(row) and row[c + 1] is not None:
                contract_number = row[c + 1]
            elif text in INCEPTION_DATE_LABELS and c + 1 < len(row) and row[c + 1] is not None:
                inception_date_value = row[c + 1]
            elif text == "fund" and header_row_idx is None:
                header_row_idx = idx
                best_date = None
                for c2, val2 in enumerate(row):
                    if val2 is None:
                        continue
                    if isinstance(val2, (datetime, date)):
                        if best_date is None or val2 > best_date:
                            best_date = val2
                            value_col = c2
                            report_date = val2
                        continue
                    val2_text = str(val2).replace("\n", " ").strip()
                    val2_lower = val2_text.lower()
                    if "initial" in val2_lower:
                        initial_contrib_col = c2
                    if "closing" in val2_lower and "market value" in val2_lower:
                        value_col = c2
                        m = DATE_IN_TEXT.search(val2_text)
                        if m:
                            parsed_d = parse_date_value(m.group(0))
                            if parsed_d:
                                report_date = parsed_d
                                best_date = parsed_d

        if header_row_idx and idx > header_row_idx and row and row[0]:
            label = str(row[0]).strip()
            label_lower = label.lower()
            if label_lower.startswith("total") and "value" in label_lower and value_col is not None:
                if "transaction" in label_lower and "progress" in label_lower:
                    total_value_including_progress = parse_money(row[value_col])
                    totals_row_idx = idx
                elif total_value is None:
                    total_value = parse_money(row[value_col])
                    if totals_row_idx is None:
                        totals_row_idx = idx
                if (total_initial_contribution is None and initial_contrib_col is not None
                        and initial_contrib_col < len(row) and row[initial_contrib_col] is not None):
                    total_initial_contribution = parse_money(row[initial_contrib_col])
            
            elif value_col is not None and value_col < len(row) and totals_row_idx is None:
                val_str = str(row[value_col])
                if val_str and label_lower not in INVESTOR_LABELS and label_lower not in PRODUCT_LABELS and label_lower not in INCEPTION_DATE_LABELS:
                    parsed_val = parse_money(val_str)
                    if parsed_val > 0:
                        funds[label] = parsed_val

    final_value = total_value_including_progress if total_value_including_progress is not None else total_value

    full_text = getattr(ws, "full_text", None)
    if not full_text:
        full_text = "\n".join(" ".join(str(v).replace("\n", " ") for v in r if v is not None) for r in rows)

    # -------------------------------------------------------------------------
    # NEW PDF FALLBACK: Explicitly target the Model Portfolio Summary Table
    # -------------------------------------------------------------------------
    if not funds and full_text:
        m_summary = re.search(r"Fund summary for the period.*?\n(.*?)\nTotal.*?[Vv]alue", full_text, re.IGNORECASE | re.DOTALL)
        if m_summary:
            for line in m_summary.group(1).split('\n'):
                line = line.strip()
                if not line or line.lower().startswith("fund"): continue
                
                m_nums = list(re.finditer(r"-?\d{1,3}(?:[,\s]\d{3})*\.\d{2}", line))
                if m_nums:
                    val = parse_money(m_nums[-1].group(0))
                    if val > 0:
                        m_name = re.search(r"^(.*?)\s+(?:\d{1,3}(?:[,\s]\d{3})*\.\d{2,4}\s+|R\s*-?\d{1,3}(?:[,\s]\d{3})*\.\d{2})", line)
                        if m_name:
                            funds[m_name.group(1).strip()] = val
    # -------------------------------------------------------------------------

    if client_name and str(client_name).strip() in ('', '|'):
        client_name = None
    if not client_name:
        first_line = full_text.strip().split('\n')[0].strip()
        if first_line and len(first_line) > 3 and not first_line.lower().startswith("allan gray"):
            client_name = first_line

    if vehicle_name and str(vehicle_name).strip() in ('', '|'):
        vehicle_name = None
    if not vehicle_name:
        m = re.search(r"Allan Gray ([A-Za-z0-9/&\-\s]+?)(?:\s+Fund)?\s+transaction history", full_text, re.IGNORECASE)
        if m:
            vehicle_name = f"Allan Gray {m.group(1).strip()}"

    if contract_number and str(contract_number).strip() in ('', '|'):
        contract_number = None
    if not contract_number:
        m = re.search(r"(?:Account|Contract|Policy)\s*number:\s*(?:\|\s*)?([A-Z0-9]+)", full_text, re.IGNORECASE)
        if m:
            contract_number = m.group(1)

    if inception_date_value and str(inception_date_value).strip() in ('', '|'):
        inception_date_value = None
    if not inception_date_value:
        m = re.search(r"(?:Inception|Commencement|Start)\s*date:\s*(?:\|\s*)?(\d{1,2}\s+[A-Za-z]+\s+\d{4})", full_text, re.IGNORECASE)
        if m:
            inception_date_value = m.group(1)

    if not report_date:
        m = re.search(r"Statement date:\s*(?:\|\s*)?(\d{1,2}\s+[A-Za-z]+\s+\d{4})", full_text, re.IGNORECASE)
        if m:
            report_date = parse_date_value(m.group(1))

    if not final_value:
        m = re.search(r"Total market value incl transactions in([\s\S]{1,150}?)(?:progress|Indicates)", full_text, re.IGNORECASE)
        if m:
            amts = re.findall(r"R\s*([0-9,\s]+\.\d{2})", m.group(1))
            if amts:
                final_value = parse_money(amts[-1])
                if not total_initial_contribution and len(amts) >= 2:
                    total_initial_contribution = parse_money(amts[0])
        if not final_value:
            m = re.search(r"Total Value([\s\S]{1,150}?)Total market", full_text, re.IGNORECASE)
            if m:
                amts = re.findall(r"R\s*([0-9,\s]+\.\d{2})", m.group(1))
                if amts:
                    final_value = parse_money(amts[-1])
                    if not total_initial_contribution and len(amts) >= 2:
                        total_initial_contribution = parse_money(amts[0])

    if not contract_number or not final_value:
        return {}

    if client_name:
        client_name = re.sub(r"\s*-\s*\d+\s*$", "", str(client_name)).strip()

    parsed_inception = parse_date_value(inception_date_value)
    contract_str = str(contract_number).strip()

    fname = os.path.basename(filepath)
    if parsed_inception is None:
        logging.info(f"{fname}: No 'Inception date' field found for account {contract_str} - left blank.")
    if total_initial_contribution is not None:
        logging.info(f"{fname}: Net Contribution set to R{total_initial_contribution:,.2f}.")
    
    return {
        contract_str: {
            "client_name": client_name or "Unknown Client",
            "vehicle_name": vehicle_name or "Allan Gray Investment",
            "value": final_value,
            "report_date": report_date,
            "transaction_date": report_date,
            "net_contribution": total_initial_contribution or 0.0,
            "inception_date": parsed_inception,
            "funds": funds,
        }
    }

class ListWorksheet:
    def __init__(self, rows, full_text=None):
        width = max((len(r) for r in rows), default=0)
        self.rows = [list(r) + [None] * (width - len(r)) for r in rows]
        self.full_text = full_text

    def iter_rows(self, min_row=1, max_row=None, max_col=None, values_only=True):
        end = max_row if max_row is not None else len(self.rows)
        for row in self.rows[min_row - 1:end]:
            yield tuple(row[:max_col] if max_col else row)

def _split_layout_line(line):
    return [c.strip() for c in re.split(r"\s{2,}", line.strip()) if c.strip() != ""]

def extract_pdf_rows_via_ocr(filepath):
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise RuntimeError("Scanned PDF detected. Install pytesseract and pdf2image.") from exc

    rows = []
    full_text_parts = []
    images = convert_from_path(filepath)
    for image in images:
        text = pytesseract.image_to_string(image)
        full_text_parts.append(text)
        for line in text.splitlines():
            cols = _split_layout_line(line)
            if cols:
                rows.append(cols)
    return rows, "\n".join(full_text_parts)

def extract_pdf_rows(filepath):
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("PDF support requires pdfplumber.") from exc

    rows = []
    full_text_parts = []
    total_text_len = 0

    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            total_text_len += len(page_text.strip())
            full_text_parts.append(page_text)
            tables = page.extract_tables()
            if tables:
                for table in tables:
                    for row in table:
                        rows.append([c.strip() if isinstance(c, str) else c for c in row])
            elif page_text.strip():
                for line in page_text.splitlines():
                    cols = _split_layout_line(line)
                    if cols:
                        rows.append(cols)

    full_text = "\n".join(full_text_parts)
    if total_text_len < 20:
        rows, full_text = extract_pdf_rows_via_ocr(filepath)

    return rows, full_text

def load_worksheet_like(filepath):
    lower = filepath.lower()
    if lower.endswith(".xlsx"):
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        return wb.active, wb
    elif lower.endswith(".pdf"):
        rows, full_text = extract_pdf_rows(filepath)
        return ListWorksheet(rows, full_text=full_text), None
    else:
        raise ValueError(f"Unsupported statement file type: {os.path.basename(filepath)}")

def read_statement(filepath):
    try:
        ws, wb = load_worksheet_like(filepath)
    except Exception as e:
        logging.error(f"Failed to open '{os.path.basename(filepath)}': {e}")
        return {}

    if is_allan_gray_fund_summary(ws):
        primary, fallback = read_allan_gray_fund_summary, read_old_mutual_style_statement
    else:
        primary, fallback = read_old_mutual_style_statement, read_allan_gray_fund_summary

    try:
        results = primary(ws, filepath)
    except Exception as exc:
        logging.warning(f"Error parsing '{os.path.basename(filepath)}' with primary parser: {exc}. Trying fallback.")
        results = {}

    if not results:
        try:
            results = fallback(ws, filepath)
        except Exception as exc:
            logging.error(f"Complete failure parsing '{os.path.basename(filepath)}': {exc}")
            results = {}

    if wb is not None:
        wb.close()
    return results

# ---------------------------------------------------------------------------
# CASH FLOW EXTRACTION FUNCTIONS
# ---------------------------------------------------------------------------

def parse_cf_pdf(file_path):
    extracted = {"account_number": None, "raw_transactions": []}
    date_pattern = re.compile(r'^(\d{2}\s[a-zA-Z]{3}\s\d{4})')
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if not text:
                    continue
                for line in text.split('\n'):
                    if not extracted["account_number"]:
                        match = re.search(r'Account number:\s*(?:\|\s*)?([A-Z0-9]+)', line, re.IGNORECASE)
                        if match: 
                            extracted["account_number"] = match.group(1).strip().upper()

                    lower_line = line.lower()
                    if any(ign in lower_line for ign in CF_IGNORE_KEYWORDS): 
                        continue
                        
                    for raw_kw, clean_desc in CF_DESCRIPTION_MAPPING.items():
                        if raw_kw in lower_line:
                            date_match = date_pattern.search(line.strip())
                            amt_match = re.search(r'(-?R\s?[\d\s,]+\.\d{2})$', line.strip())
                            
                            if date_match and amt_match:
                                raw_amt = amt_match.group(1).replace(" ", "").replace("R", "").replace(",", "")
                                try:
                                    extracted["raw_transactions"].append({
                                        "DateObj": datetime.strptime(date_match.group(1), "%d %b %Y"),
                                        "Description": clean_desc,
                                        "Amount": float(raw_amt)
                                    })
                                except ValueError: 
                                    pass
                                break 
    except Exception as e:
        logging.error(f"CF PDF read failed for {file_path}: {e}")
    return extracted

def parse_cf_excel(file_path):
    extracted = {"account_number": None, "raw_transactions": []}
    date_pattern = re.compile(r'^(\d{2}\s[a-zA-Z]{3}\s\d{4})')
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active 
        
        for row in ws.iter_rows(values_only=True):
            row_clean = [c for c in row if c is not None and str(c).strip() != ""]
            if not row_clean: 
                continue
            
            row_str = " ".join(str(c) for c in row_clean).lower()
            if not extracted["account_number"]:
                match = re.search(r'account number[\s:|]*([a-z0-9]+)', row_str)
                if match: 
                    extracted["account_number"] = match.group(1).upper()

            if any(ign in row_str for ign in CF_IGNORE_KEYWORDS): 
                continue
            
            for raw_kw, clean_desc in CF_DESCRIPTION_MAPPING.items():
                if raw_kw in row_str:
                    date_obj = None
                    amount_val = None
                    for cell in row_clean:
                        if isinstance(cell, datetime):
                            date_obj = cell
                        elif isinstance(cell, str) and date_pattern.search(cell.strip()):
                            try:
                                date_obj = datetime.strptime(date_pattern.search(cell.strip()).group(1), "%d %b %Y")
                            except ValueError:
                                pass
                        
                        if isinstance(cell, (int, float)):
                            amount_val = float(cell)
                        elif isinstance(cell, str):
                            amt_match = re.search(r'(-?R\s?[\d\s,]+\.\d{2})', cell)
                            if amt_match:
                                raw_amt = amt_match.group(1).replace(" ", "").replace("R", "").replace(",", "")
                                try:
                                    amount_val = float(raw_amt)
                                except ValueError:
                                    pass
                    
                    if date_obj and amount_val is not None:
                        extracted["raw_transactions"].append({
                            "DateObj": date_obj, 
                            "Description": clean_desc, 
                            "Amount": amount_val
                        })
                        break
        wb.close()
    except Exception as e:
        logging.error(f"CF Excel read failed for {file_path}: {e}")
    return extracted

def consolidate_cf_transactions(account_num, raw_transactions):
    if not raw_transactions: 
        return []
        
    daily_consolidated = defaultdict(float)
    for t in raw_transactions: 
        daily_consolidated[(t["DateObj"], t["Description"])] += t["Amount"]

    date_groups = defaultdict(list)
    for (d_obj, desc), amt in daily_consolidated.items():
        if abs(amt) > 0.01: 
            date_groups[d_obj].append({"Description": desc, "Amount": amt})

    consolidated_list = []
    for d_obj, flows in date_groups.items():
        pos_flows = [f for f in flows if f["Amount"] > 0]
        neg_flows = [f for f in flows if f["Amount"] < 0]
        used_pos = set()
        used_neg = set()
        
        for p_idx, p in enumerate(pos_flows):
            for n_idx, n in enumerate(neg_flows):
                if n_idx not in used_neg and abs(p["Amount"] + n["Amount"]) < 0.01:
                    used_pos.add(p_idx)
                    used_neg.add(n_idx)
                    break
        
        for p_idx, p in enumerate(pos_flows):
            if p_idx not in used_pos: 
                consolidated_list.append({"DateObj": d_obj, "Description": p["Description"], "Amount": p["Amount"]})
        for n_idx, n in enumerate(neg_flows):
            if n_idx not in used_neg: 
                consolidated_list.append({"DateObj": d_obj, "Description": n["Description"], "Amount": n["Amount"]})

    consolidated_list.sort(key=lambda x: x["DateObj"])
    if not consolidated_list: 
        return []

    prefix = account_num[:4]
    
    initial_inv = None
    grouped_inflows = []
    grouped_outflows = []
    itemized_outflows = []

    for cf in consolidated_list:
        if cf["Amount"] > 0:
            if not initial_inv:
                initial_inv = cf
            else:
                grouped_inflows.append(cf)
        else:
            if prefix == "AGLA":
                grouped_outflows.append(cf)
            else:
                itemized_outflows.append(cf)

    final_flows = []

    if initial_inv:
        final_flows.append({
            "DateStr": initial_inv["DateObj"].strftime("%d/%m/%Y"),
            "Description": "Initial Investment",
            "Amount": initial_inv["Amount"]
        })

    if grouped_inflows:
        min_date = min(f["DateObj"] for f in grouped_inflows)
        max_date = max(f["DateObj"] for f in grouped_inflows)
        total_amt = sum(f["Amount"] for f in grouped_inflows)
        
        if min_date == max_date:
            date_str = min_date.strftime('%d/%m/%Y')
        else:
            date_str = f"{min_date.strftime('%d/%m/%Y')} - {max_date.strftime('%d/%m/%Y')}"
        
        if prefix == "AGTF":
            desc = "Total Debit Orders"
            if (datetime.now() - max_date).days > 45:
                desc += f" - Stopped {max_date.strftime('%B %Y')}"
        else:
            desc = "Total Additional Contributions"

        final_flows.append({
            "DateStr": date_str,
            "Description": desc,
            "Amount": total_amt
        })

    if grouped_outflows:
        min_date = min(f["DateObj"] for f in grouped_outflows)
        max_date = max(f["DateObj"] for f in grouped_outflows)
        total_amt = sum(f["Amount"] for f in grouped_outflows)
        
        if min_date == max_date:
            date_str = min_date.strftime('%d/%m/%Y')
        else:
            date_str = f"{min_date.strftime('%d/%m/%Y')} - {max_date.strftime('%d/%m/%Y')}"
        
        final_flows.append({
            "DateStr": date_str,
            "Description": "Total Annuity Income Payments",
            "Amount": total_amt
        })

    for cf in itemized_outflows:
        final_flows.append({
            "DateStr": cf["DateObj"].strftime("%d/%m/%Y"),
            "Description": cf["Description"],
            "Amount": cf["Amount"]
        })
        
    return final_flows

def build_cashflow_tab(wb, accounts, date_tab_str):
    target_sheet_name = next((s for s in wb.sheetnames if s.replace(" ", "").lower().startswith("cashflow")), None)
    if not target_sheet_name:
        logging.warning("No 'Cash Flows' template tab found in workbook. Skipping Cash Flow generation.")
        return

    # Update sheet IN PLACE so images aren't deleted by copy operations
    sheet = wb[target_sheet_name]
    sheet.title = f"CashFlows_{date_tab_str}"

    for m_range in list(sheet.merged_cells.ranges): 
        sheet.unmerge_cells(str(m_range))

    grouped_data = defaultdict(list)
    for acc_num, info in accounts.items():
        cfs = info.get("cash_flows")
        if not cfs: 
            continue
        cat = CF_ACCOUNT_MAPPING.get(acc_num[:4], {}).get("category", "Discretionary Investments")
        grouped_data[cat].append({"account_number": acc_num, "cash_flows": cfs})

    if not grouped_data:
        return

    col_map = {"account": 1, "provider": 3, "product": 5, "date": 7, "desc": 9, "amount": 13}
    header_row = 4
    for r in range(1, 10):
        val = str(sheet.cell(row=r, column=1).value or "").strip().lower()
        if "account number" in val:
            header_row = r
            break

    if sheet.max_row > header_row:
        sheet.delete_rows(header_row + 1, sheet.max_row - header_row)

    current_row = header_row + 1
    
    standard_font = Font(name='Arial', size=10, color="000000")
    bold_font = Font(name='Arial', size=10, bold=True, color="000000")
    category_font = Font(name='Arial', size=11, bold=True, italic=True, color="000000")
    accounting_border = Border(top=Side(style='thin'), bottom=Side(style='double'))
    financial_format = '[$R-en-ZA] #,##0.00;[Red]-[$R-en-ZA] #,##0.00'

    for cat_name in CF_CATEGORY_ORDER:
        if cat_name in grouped_data:
            c_cat = sheet.cell(row=current_row, column=col_map["account"])
            c_cat.value = cat_name
            c_cat.font = category_font
            current_row += 1
            
            for acc_data in grouped_data[cat_name]:
                acc_num = acc_data["account_number"]
                cash_flows = acc_data["cash_flows"]
                product = CF_ACCOUNT_MAPPING.get(acc_num[:4], {}).get("product", "Unknown Product")
                
                start_block_row = current_row

                for i, cf in enumerate(cash_flows):
                    sheet.row_dimensions[current_row].height = 15
                    
                    if i == 0:
                        sheet.cell(row=current_row, column=col_map["account"]).value = f"({acc_num})"
                        sheet.cell(row=current_row, column=col_map["provider"]).value = "Allan Gray"
                        sheet.cell(row=current_row, column=col_map["product"]).value = product

                    sheet.cell(row=current_row, column=col_map["date"]).value = cf["DateStr"]
                    sheet.cell(row=current_row, column=col_map["desc"]).value = cf["Description"]
                    
                    c_amt = sheet.cell(row=current_row, column=col_map["amount"])
                    c_amt.value = cf["Amount"]
                    c_amt.number_format = financial_format
                    
                    for c in col_map.values():
                        cell = sheet.cell(row=current_row, column=c)
                        cell.font = standard_font
                        align = 'right' if c == col_map["amount"] else 'left'
                        cell.alignment = Alignment(vertical='bottom', horizontal=align, wrap_text=False)
                    
                    current_row += 1

                sheet.row_dimensions[current_row].height = 15
                
                c_lbl = sheet.cell(row=current_row, column=col_map["desc"])
                c_lbl.value = "Net Contributions"
                c_lbl.font = bold_font
                c_lbl.alignment = Alignment(vertical='bottom', horizontal='right')

                tot_cell = sheet.cell(row=current_row, column=col_map["amount"])
                amt_letter = get_column_letter(col_map["amount"])
                
                tot_cell.value = f"=SUM({amt_letter}{start_block_row}:{amt_letter}{current_row - 1})"
                tot_cell.font = bold_font
                tot_cell.border = accounting_border
                tot_cell.number_format = financial_format
                tot_cell.alignment = Alignment(vertical='bottom', horizontal='right', wrap_text=False)
                
                current_row += 2 

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=14)

# ---------------------------------------------------------------------------
# MAIN SCRIPT LOGIC
# ---------------------------------------------------------------------------

def match_row(vehicle_name):
    text = (vehicle_name or "").lower()
    for row_num, (label, keywords, excludes) in PLACEHOLDER_ROWS.items():
        if any(k in text for k in keywords) and not any(e in text for e in excludes):
            return row_num, label
    return None, None

def safe_filename_part(text):
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    return text.strip("_")

def safe_sheet_name(base_name, existing_titles):
    name = re.sub(r"[\[\]:\*\?/\\]", "_", base_name)[:31]
    if name not in existing_titles:
        return name
    suffix = 1
    while True:
        candidate = f"{name[:31 - len(str(suffix)) - 1]}_{suffix}"
        if candidate not in existing_titles:
            return candidate
        suffix += 1

def clear_placeholder_spaces(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == " ":
                cell.value = None

def check_template_is_current(ws, template_file):
    b9 = ws.cell(row=9, column=INVESTMENTS_COLUMN)
    if isinstance(b9, MergedCell) or (str(ws.cell(row=9, column=1).value or "")).strip().startswith("Sub Total"):
        print(f"\nERROR: '{os.path.basename(template_file)}' looks like an OUTDATED template.")
        return False
    return True

def get_latest_transaction_date(accounts):
    latest = None
    for info in accounts.values():
        parsed = parse_date_value(info.get("transaction_date"))
        if parsed and (latest is None or parsed > latest):
            latest = parsed
    return latest

def parse_ddmmyy_sheet_name(name):
    if not re.fullmatch(r"\d{6}", name):
        return None
    try:
        return datetime.strptime(name, "%d%m%y").date()
    except ValueError:
        return None

def find_latest_dated_tab(wb):
    latest_date, latest_ws = None, None
    for sheet_name in wb.sheetnames:
        sheet_date = parse_ddmmyy_sheet_name(sheet_name)
        if sheet_date and (latest_date is None or sheet_date > latest_date):
            latest_date, latest_ws = sheet_date, wb[sheet_name]
    return latest_ws

def print_data_extraction_preview(by_client):
    print("=" * 80)
    print("DATA EXTRACTION PREVIEW".center(80))
    print("=" * 80)
    
    for client_name, accounts in by_client.items():
        print(f"\nCLIENT: {client_name}")
        for contract, info in accounts.items():
            print(f"  Account: {contract} ({info.get('vehicle_name', 'Unknown')})")
            print(f"    Total Value:      R {info.get('value', 0.0):,.2f}")
            print(f"    Net Contribution: R {info.get('net_contribution', 0.0):,.2f}")
            
            funds = info.get("funds", {})
            if funds:
                print("    Funds Breakdown:")
                for fund_name, fund_val in funds.items():
                    print(f"      - {fund_name}: R {fund_val:,.2f}")
            else:
                print("    Funds Breakdown: [NO UNDERLYING FUNDS FOUND]")
    print("\n" + "=" * 80 + "\n")


def fill_client_workbook(client_name, accounts, summary_file):
    wb = openpyxl.load_workbook(summary_file)

    try:
        source_ws = wb[TEMPLATE_SHEET_NAME]
    except KeyError:
        logging.error(f"Could not find target template sheet '{TEMPLATE_SHEET_NAME}' in {summary_file}")
        return None

    if not check_template_is_current(source_ws, summary_file):
        return None

    latest_txn_date = get_latest_transaction_date(accounts)
    used_fallback_date = latest_txn_date is None
    if used_fallback_date:
        latest_txn_date = date.today()

    date_title_str = latest_txn_date.strftime("%d %B %Y")
    date_tab_str = latest_txn_date.strftime("%d%m%y")

    # -------------------------------------------------------------------------
    # TAB ORDERING FIX: Place directly right of the latest transaction tab
    # -------------------------------------------------------------------------
    latest_dated_ws = find_latest_dated_tab(wb)
    new_sheet_name = safe_sheet_name(date_tab_str, set(wb.sheetnames))
    ws = wb.copy_worksheet(source_ws)
    ws.title = new_sheet_name

    sheets = wb._sheets
    sheets.remove(ws) 
    
    if latest_dated_ws:
        latest_idx = sheets.index(latest_dated_ws)
        sheets.insert(latest_idx + 1, ws)
    else:
        source_idx = sheets.index(source_ws)
        sheets.insert(source_idx + 1, ws)
    # -------------------------------------------------------------------------

    if getattr(source_ws, '_images', []):
        try:
            with zipfile.ZipFile(summary_file, 'r') as zf:
                for img in source_ws._images:
                    clean_path = img.path[1:] if img.path.startswith('/') else img.path
                    if clean_path in zf.namelist():
                        img_data = zf.read(clean_path)
                        new_img = OpenpyxlImage(io.BytesIO(img_data))
                        new_img.anchor = copy.deepcopy(img.anchor) 
                        ws.add_image(new_img)
        except Exception as e:
            logging.warning(f"Could not copy logo/image to new tab: {e}")

    clear_placeholder_spaces(ws)

    # -------------------------------------------------------------------------
    # UNMERGE EVERYTHING TO PREVENT LAYOUT BUGS
    # -------------------------------------------------------------------------
    ranges_to_unmerge = []
    for m_range in ws.merged_cells.ranges:
        ranges_to_unmerge.append(str(m_range))
            
    for r_str in ranges_to_unmerge:
        ws.unmerge_cells(r_str)
    # -------------------------------------------------------------------------

    # -------------------------------------------------------------------------
    # WIPE ALL LINGERING "(Investment Strategy)" GHOST TEXT
    # -------------------------------------------------------------------------
    for r in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=r, column=INVESTMENTS_COLUMN).value
        if cell_val and isinstance(cell_val, str) and cell_val.strip().startswith("(Investment Strategy)"):
            ws.cell(row=r, column=INVESTMENTS_COLUMN).value = None
    # -------------------------------------------------------------------------

    title_cell = ws.cell(row=TITLE_ROW, column=TITLE_COLUMN)
    if title_cell.value:
        new_title = str(title_cell.value)
        if "Client Name & Surname" in new_title or "DD/Month/YYYY" in new_title:
            new_title = re.sub(r"Client Name & Surname", client_name, new_title)
            new_title = re.sub(r"DD/Month/YYYY", date_title_str, new_title)
        else:
            month_pattern = (r"\b\d{1,2}\s+(?:January|February|March|April|May|June|July|"
                              r"August|September|October|November|December)\s+\d{4}\b")
            if re.search(month_pattern, new_title):
                new_title = re.sub(month_pattern, date_title_str, new_title)
        
        ws.cell(row=TITLE_ROW, column=TITLE_COLUMN).value = new_title
        ws.cell(row=TITLE_ROW, column=TITLE_COLUMN).font = Font(name='Arial', size=29, bold=True, color="000000")

    row_allocations = {}
    unmatched = []

    for contract, info in accounts.items():
        row_num, label = match_row(info["vehicle_name"])
        if row_num is None:
            unmatched.append((contract, info))
        else:
            row_allocations.setdefault(row_num, []).append((contract, info, label))

    category_totals = {
        "Retirement": {"val": 0.0, "contrib": 0.0},
        "Discretionary": {"val": 0.0, "contrib": 0.0},
        "Offshore": {"val": 0.0, "contrib": 0.0}
    }

    matched = []
    
    dotted_side = Side(border_style="dotted", color="000000")
    dotted_border = Border(left=dotted_side, right=dotted_side, top=dotted_side, bottom=dotted_side)
    
    actions = []
    for ph_row in PLACEHOLDER_ROWS.keys():
        if ph_row in row_allocations:
            actions.append(('fill', ph_row))
        else:
            actions.append(('delete', ph_row))
            
    # CRITICAL: Process backwards so deleting/inserting rows doesn't break the rows above them
    actions.sort(key=lambda x: x[1], reverse=True)
    
    for action, row_num in actions:
        span = PLACEHOLDER_SPANS.get(row_num, 2)
        
        if action == 'delete':
            ws.delete_rows(row_num, amount=span)
            continue
            
        acc_list = row_allocations[row_num]
        
        total_value = sum(a[1]['value'] for a in acc_list)
        total_net_contribution = sum(a[1].get('net_contribution', 0.0) for a in acc_list)
        inception_dates_found = [a[1]['inception_date'] for a in acc_list if a[1].get('inception_date')]
        combined_inception_date = min(inception_dates_found) if inception_dates_found else None
        label = acc_list[0][2]
        
        if row_num in [5, 7, 9]:
            category_totals["Retirement"]["val"] += total_value
            category_totals["Retirement"]["contrib"] += total_net_contribution
        elif row_num in [13, 16]:
            category_totals["Discretionary"]["val"] += total_value
            category_totals["Discretionary"]["contrib"] += total_net_contribution
        elif row_num in [21, 23]:
            category_totals["Offshore"]["val"] += total_value
            category_totals["Offshore"]["contrib"] += total_net_contribution

        lines_to_write = []
        for idx, (contract, info, _) in enumerate(acc_list):
            vehicle = str(info.get('vehicle_name', 'Unknown'))
            
            # STRIP CLASSES AND ASTERISKS FROM MAIN VEHICLE NAME
            vehicle = re.sub(r"(?i)\s*\(?Class\s+[A-Z0-9]+\)?", "", vehicle).strip()
            vehicle = vehicle.replace("**", "").replace("*", "").strip()
            
            if vehicle.lower().startswith("allan gray"):
                provider = "Allan Gray"
                product = vehicle[10:].strip()
            elif vehicle.lower().startswith("momentum"):
                provider = "Momentum"
                product = vehicle[8:].strip()
            else:
                provider = "Service Provider"
                product = vehicle

            account_name = f"{provider} {product} - {contract}"
            
            lines_to_write.append({
                'type': 'main',
                'name': account_name,
                'value': info['value'],
                'is_first_of_category': (idx == 0)
            })
            
            for fund_name, fund_val in info.get("funds", {}).items():
                # STRIP CLASSES AND ASTERISKS FROM UNDERLYING FUND NAMES
                clean_fund = re.sub(r"(?i)\s*\(?Class\s+[A-Z0-9]+\)?", "", str(fund_name)).strip()
                clean_fund = clean_fund.replace("**", "").replace("*", "").strip()
                
                lines_to_write.append({
                    'type': 'fund',
                    'name': clean_fund,
                    'value': fund_val
                })
            
            # Append a blank spacing row AFTER every account block
            lines_to_write.append({
                'type': 'blank',
                'name': '',
                'value': None
            })

        rows_needed = len(lines_to_write)
        rows_to_insert = rows_needed - span

        # --- PRESERVE COMMENTS (COLUMNS F TO O) ---
        preserved_comments = {}
        for col_c in range(6, 16):
            preserved_comments[col_c] = ws.cell(row=row_num, column=col_c).value
        # ------------------------------------------

        if rows_to_insert > 0:
            ws.insert_rows(row_num + span, amount=rows_to_insert)
        elif rows_to_insert < 0:
            ws.delete_rows(row_num + rows_needed, amount=abs(rows_to_insert))

        for i, line in enumerate(lines_to_write):
            curr_r = row_num + i
            
            # Wipe ONLY columns 1 to 5 to protect original side comments. 
            for c in range(1, 6):
                ws.cell(row=curr_r, column=c).value = None
            
            # If it's the very first row of the section, paste the preserved comment text back
            if i == 0:
                for col_c, val in preserved_comments.items():
                    ws.cell(row=curr_r, column=col_c).value = val
            else:
                # Clear comments for any subsequent lines to avoid repeating text down the page
                for c in range(6, 16):
                    ws.cell(row=curr_r, column=c).value = None
            
            # If it's a spacing row, just paint the dotted borders and continue
            if line['type'] == 'blank':
                for col in range(1, 6):
                    ws.cell(row=curr_r, column=col).border = dotted_border
                ws.row_dimensions[curr_r].height = None
                continue
            
            cell_name = ws.cell(row=curr_r, column=INVESTMENTS_COLUMN)
            cell_name.value = line['name']
            
            cell_name.alignment = Alignment(wrap_text=True, horizontal='left', vertical='center')
            
            cell_val = ws.cell(row=curr_r, column=CURRENT_VALUE_COLUMN)
            if line['value'] is not None and line['value'] != 0.0:
                cell_val.value = line['value']
                cell_val.number_format = '[$R-en-ZA] #,##0.00'
            cell_val.alignment = Alignment(wrap_text=True, horizontal='right', vertical='center')
            
            if line['type'] == 'main':
                cell_name.font = Font(name='Arial', size=29, bold=True, color="000000")
                cell_val.font = Font(name='Arial', size=29, bold=True, color="000000")
                
                for col in range(1, 6):
                    ws.cell(row=curr_r, column=col).border = dotted_border
                
                ws.row_dimensions[curr_r].height = None
                
                if line.get('is_first_of_category'):
                    name_cell = ws.cell(row=curr_r, column=CLIENT_NAME_COLUMN)
                    name_cell.value = client_name
                    name_cell.font = Font(name='Arial', size=29, bold=True, color="000000")
                    name_cell.alignment = Alignment(wrap_text=True, horizontal='left', vertical='center')
                    
                    if combined_inception_date:
                        inc_cell = ws.cell(row=curr_r, column=INCEPTION_DATE_COLUMN)
                        inc_cell.value = combined_inception_date
                        inc_cell.number_format = 'DD/MM/YYYY'
                        inc_cell.font = Font(name='Arial', size=29, color="000000")
                        inc_cell.alignment = Alignment(horizontal='center', vertical='center')
                        
                    if total_net_contribution:
                        contrib_cell = ws.cell(row=curr_r, column=NET_CONTRIB_COLUMN)
                        contrib_cell.value = total_net_contribution
                        contrib_cell.number_format = '[$R-en-ZA] #,##0.00'
                        contrib_cell.font = Font(name='Arial', size=29, bold=True, color="000000")
                        contrib_cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell_name.font = Font(name='Arial', size=29, italic=True, color="000000")
                cell_val.font = Font(name='Arial', size=29, bold=False, color="000000")
                
                for col in range(1, 6):
                    ws.cell(row=curr_r, column=col).border = dotted_border

                ws.row_dimensions[curr_r].height = None

        matched.append((f"{acc_list[0][1]['vehicle_name']}...", total_value, row_num, label))

    grand_total_val = 0.0
    grand_total_contrib = 0.0
    
    for r in range(1, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val and isinstance(val, str):
            text = val.strip().lower()
            
            is_subtotal = False
            
            if "sub total retirement" in text:
                cat_val = category_totals["Retirement"]["val"]
                cat_contrib = category_totals["Retirement"]["contrib"]
                is_subtotal = True
                grand_total_val += cat_val
                grand_total_contrib += cat_contrib
            elif "sub total discretionary investments - offshore" in text:
                cat_val = category_totals["Offshore"]["val"]
                cat_contrib = category_totals["Offshore"]["contrib"]
                is_subtotal = True
                grand_total_val += cat_val
                grand_total_contrib += cat_contrib
            elif "sub total discretionary" in text: 
                cat_val = category_totals["Discretionary"]["val"]
                cat_contrib = category_totals["Discretionary"]["contrib"]
                is_subtotal = True
                grand_total_val += cat_val
                grand_total_contrib += cat_contrib
            elif text == "total":
                cat_val = grand_total_val
                cat_contrib = grand_total_contrib
                is_subtotal = True
                
            if is_subtotal:
                c_val = ws.cell(row=r, column=5)
                c_val.value = cat_val
                c_val.number_format = '[$R-en-ZA] #,##0.00'
                c_val.font = Font(name='Arial', size=29, bold=True, color="000000")
                
                c_contrib = ws.cell(row=r, column=4)
                c_contrib.value = cat_contrib
                c_contrib.number_format = '[$R-en-ZA] #,##0.00'
                c_contrib.font = Font(name='Arial', size=29, bold=True, color="000000")

    # -------------------------------------------------------------------------
    # HEADER PADDING: Insert 1 blank row directly under Headers
    # -------------------------------------------------------------------------
    headers = []
    for r in range(1, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val and isinstance(val, str):
            text = val.strip()
            if text in ["Retirement Funding", "Discretionary Investments - Local", "Discretionary Investments - Offshore", "Cash Flows"]:
                headers.append(r)
                
    headers.sort(reverse=True)
    for r in headers:
        ws.insert_rows(r + 1, 1)
        # We ensure it's unbordered so it looks like clean padding
        for c in range(1, 11):
            ws.cell(row=r + 1, column=c).border = Border()
    # -------------------------------------------------------------------------

    # -------------------------------------------------------------------------
    # RESTORE FORMATTING: Merge & Center / Wrap Text ONLY for Main Headings 
    # -------------------------------------------------------------------------
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=10)
    ws.cell(row=TITLE_ROW, column=1).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for r in range(1, ws.max_row + 1):
        val = ws.cell(row=r, column=1).value
        if val and isinstance(val, str):
            text = val.strip()
            # Main Section Headings (A to J)
            if text in ["Retirement Funding", "Discretionary Investments - Local", "Discretionary Investments - Offshore", "Cash Flows"]:
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
                ws.cell(row=r, column=1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            
            # Sub Totals (A to B only)
            elif text.lower().startswith("sub total") or text.lower() == "total":
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
                ws.cell(row=r, column=1).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    # -------------------------------------------------------------------------

    return wb, matched, unmatched, date_tab_str, used_fallback_date, source_ws.title


def main():
    print(f"Looking in: {FOLDER}\n")
    extract_zip_files()

    summary_files = find_investment_summary_files()
    if not summary_files:
        logging.error("Could not find any existing Investment Summary file in this folder.")
        return
    
    statement_files = find_statement_files(summary_files)
    if not statement_files:
        print("No .xlsx or .pdf statement files found. Nothing to do.")
        return

    all_accounts = {}
    all_cf_data = {}

    print("=" * 80)
    print("PHASE 1: DUAL DATA EXTRACTION".center(80))
    print("=" * 80)

    for f in statement_files:
        print(f"[LOG] Extracting: {os.path.basename(f)}")
        
        # 1. Main Dashboard Extraction
        accounts = read_statement(f)
        for contract, info in accounts.items():
            if contract not in all_accounts:
                all_accounts[contract] = info
                
        # 2. Cash Flow Extraction
        if f.lower().endswith('.pdf'):
            cf_ext = parse_cf_pdf(f)
        else:
            cf_ext = parse_cf_excel(f)
            
        if cf_ext.get("account_number") and cf_ext.get("raw_transactions"):
            cf_ext["cash_flows"] = consolidate_cf_transactions(cf_ext["account_number"], cf_ext["raw_transactions"])
            all_cf_data[cf_ext["account_number"]] = cf_ext

    if not all_accounts:
        print("\nNo valid statement data found.")
        return

    # THE DATA HANDSHAKE: Link Cash Flows to Client Accounts
    for contract, info in all_accounts.items():
        if contract in all_cf_data:
            cfs = all_cf_data[contract]["cash_flows"]
            # Override the dashboard net contribution with the mathematically verified CF total
            info["net_contribution"] = sum(cf["Amount"] for cf in cfs)
            info["cash_flows"] = cfs
        else:
            info["cash_flows"] = []

    by_client = {}
    for contract, info in all_accounts.items():
        norm_name = normalize_client_name(info["client_name"])
        info["client_name"] = norm_name
        by_client.setdefault(norm_name, {})[contract] = info

    print_data_extraction_preview(by_client)

    print(f"Building summaries for {len(by_client)} client(s):\n")
    for client_name, accounts in by_client.items():
        
        # Determine the target file; fall back to generic template if client is new
        summary_file = match_summary_file_for_client(client_name, summary_files)
        if summary_file is None:
            template_files = [f for f in summary_files if "template" in normalize(os.path.basename(f))]
            if template_files:
                summary_file = template_files[0]
                logging.info(f"No existing summary for {client_name}. Using Template.")
            else:
                summary_file = summary_files[0]
                logging.info(f"No existing summary or explicit template for {client_name}. Using first available file as base.")

        result = fill_client_workbook(client_name, accounts, summary_file)

        if result is None:
            continue

        wb, matched, unmatched, date_tab_str, used_fallback_date, source_tab_name = result
        
        # INJECT CASH FLOW TAB (Updates in place to protect images)
        build_cashflow_tab(wb, accounts, date_tab_str)

        filename = f"{OUTPUT_PREFIX}{safe_filename_part(client_name)}_{date_tab_str}.xlsx"
        output_path = os.path.join(FOLDER, filename)

        try:
            wb.save(output_path)
            print(f" [SUCCESS] Saved to {filename}")
        except PermissionError:
            print(f"\n[ERROR] Could not save '{filename}'. Please close it in Excel if open.")
            continue

        for match in matched:
            print(f"    - Row {match[2]} ({match[3]}): R{match[1]:,.2f}")
            
        # -------------------------------------------------------------------------
        # TERMINAL WARNING SYSTEM: Flag Missing Values before completing
        # -------------------------------------------------------------------------
        missing_reports = []
        for contract, info in accounts.items():
            missing = []
            if not info.get('inception_date'):
                missing.append("Inception Date")
            if not info.get('net_contribution'):
                missing.append("Net Contribution")
            if not info.get('funds'):
                missing.append("Fund Breakdown (Empty)")
            
            if missing:
                prod_name = info.get('vehicle_name', 'Unknown Product')
                missing_reports.append(f"      * {prod_name} ({contract}): Missing {', '.join(missing)}")
        
        if missing_reports:
            print("\n    [!] MISSING DATA / ATTENTION REQUIRED:")
            for r in missing_reports:
                print(r)
        # -------------------------------------------------------------------------
        
        print()

    print("Done.")

if __name__ == "__main__":
    main()
"""
=== AI EXECUTION INSTRUCTIONS (READ FIRST) ===
Dear AI Assistant: 
The user wants you to run this script to process their uploaded Excel files. 
Please follow these exact steps:
1. Ensure 'openpyxl' is available in your Python environment.
2. Place all uploaded files (the template and the statements) in your current working directory.
3. Execute this entire script as-is.
4. The script will generate new Excel files starting with "Investment_Summary_".
5. CRITICAL: You must return the newly generated Excel (.xlsx) files directly into this chat so the user can download them. If there are multiple files, ZIP them together into a single file and provide the download button/link.
6. Print the console output of the script to the user so they can see what matched and what was skipped.
==============================================

Veritas Work - Investment Summary Builder (AI-Ready Version)
"""

import os
import re
import zipfile
import logging
from datetime import date, datetime

import openpyxl
from openpyxl.cell.cell import MergedCell

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------
logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')

# AI Environment Fix: Jupyter notebooks (used by AI) don't always have __file__ defined.
try:
    FOLDER = os.path.dirname(os.path.abspath(__file__))
except NameError:
    FOLDER = os.getcwd()

OUTPUT_PREFIX = "Investment_Summary_"
TEMPLATE_SHEET_NAME = "DDMMYY"
CLOSING_BALANCE_LABEL = "Closing Balance"

CLIENT_NAME_COLUMN = 1   # column A
INVESTMENTS_COLUMN = 2   # column B
NET_CONTRIB_COLUMN = 4   # column D
CURRENT_VALUE_COLUMN = 5  # column E
TITLE_ROW = 2
TITLE_COLUMN = 1

COL_REPORT_DATE = 0
COL_CLIENT_NAME = 4
COL_CONTRACT_NUMBER = 5
COL_VEHICLE_NAME = 7
COL_HOLDING = 9
COL_TRANSACTION_DATE = 10
COL_TRANSACTION_TYPE = 11
COL_TRANSACTION_AMOUNT = 14

PLACEHOLDER_ROWS = {
    5:  ("Living Annuity", ["living annuity"], []),
    7:  ("Retirement Annuity", ["retirement annuity"], ["living", "provident", "preservation"]),
    9:  ("Provident/Preservation Fund", ["provident", "preservation"], []),
    13: ("Investment Platform Unit Trust", ["unit trust", "investment platform"], []),
    16: ("Tax-Free Investment", ["tax-free", "tax free"], []),
    21: ("Offshore Investment", ["offshore"], ["momentum", "international"]),
    23: ("Momentum Wealth International", ["momentum", "international"], []),
}
VALUE_WRITE_ROW_OVERRIDE = {13: 14} 

COLUMN_HEADER_ALIASES = {
    "report_date": ["report date", "statement date", "valuation date"],
    "client_name": ["client name", "investor name", "member name"],
    "contract_number": ["investment contract number", "contract number",
                        "policy number", "account number"],
    "vehicle_name": ["investment vehicle name", "product name",
                     "investment name", "product"],
    "holding": ["holdings", "holding", "fund name", "fund"],
    "transaction_date": ["transaction date"],
    "transaction_type": ["transaction type"],
    "transaction_amount": ["transaction amount", "amount"],
}

FALLBACK_POSITIONS = {
    "report_date": COL_REPORT_DATE,
    "client_name": COL_CLIENT_NAME,
    "contract_number": COL_CONTRACT_NUMBER,
    "vehicle_name": COL_VEHICLE_NAME,
    "holding": COL_HOLDING,
    "transaction_date": COL_TRANSACTION_DATE,
    "transaction_type": COL_TRANSACTION_TYPE,
    "transaction_amount": COL_TRANSACTION_AMOUNT,
}

# ---------------------------------------------------------------------------

def extract_zip_files():
    """Extract zip files handling nested directories."""
    extracted_any = False
    for root, dirs, files in os.walk(FOLDER):
        for name in files:
            if not name.lower().endswith(".zip"):
                continue
            zip_path = os.path.join(root, name)
            try:
                with zipfile.ZipFile(zip_path) as zf:
                    zf.extractall(FOLDER)
                logging.info(f"Extracted zip file: {name}")
                extracted_any = True
            except Exception as exc:
                logging.warning(f"Could not extract '{name}': {exc}")
    return extracted_any

def normalize(name):
    return re.sub(r"[\s_\-]+", "", name.lower())

def normalize_client_name(name):
    """Normalize client names to group similar names (e.g. Mr J Doe -> J Doe)."""
    if not name: return "Unknown Client"
    name = str(name).strip().upper()
    name = re.sub(r"^(MR|MRS|MS|DR|PROF)\.?\s+", "", name)
    name = re.sub(r"\s+", " ", name)
    return name

def find_template_file():
    candidates = []
    for root, dirs, files in os.walk(FOLDER):
        for name in files:
            if not name.lower().endswith(".xlsx"):
                continue
            if "investmentsummarytemplate" in normalize(name) and not name.startswith("~$"):
                candidates.append(os.path.join(root, name))
    if not candidates:
        return None
    return candidates[0]

def find_statement_files(template_file):
    files = []
    for root, dirs, files_list in os.walk(FOLDER):
        for name in files_list:
            if not name.lower().endswith(".xlsx") or name.startswith("~$") or name.startswith(OUTPUT_PREFIX):
                continue
            full_path = os.path.join(root, name)
            if template_file and os.path.abspath(full_path) == os.path.abspath(template_file):
                continue
            files.append(full_path)
    return files

def parse_money(value):
    """Robustly parse money, stripping non-breaking spaces and handling (100.00)."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace('\xa0', ' ')
    text = text.replace("R", "").replace(",", "").replace(" ", "")
    if text in ("", "-"):
        return 0.0
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return float(text)
    except ValueError:
        return 0.0

def build_column_map(ws):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    headers = [str(h).strip().lower() if h is not None else "" for h in header_row]

    column_map = {}
    warnings = []
    for role, aliases in COLUMN_HEADER_ALIASES.items():
        found_index = None
        for alias in aliases:
            for i, h in enumerate(headers):
                if alias in h:
                    found_index = i
                    break
            if found_index is not None:
                break
        if found_index is None:
            found_index = FALLBACK_POSITIONS[role]
            warnings.append(f"Header '{role}' not found - falling back to col {found_index + 1}.")
        column_map[role] = found_index
    return column_map, warnings

def is_allan_gray_fund_summary(ws):
    try:
        for row in ws.iter_rows(min_row=1, max_row=20, max_col=6, values_only=True):
            for val in row:
                if val is not None and str(val).strip().lower() == "investor":
                    return True
    except Exception:
        pass
    return False

def read_old_mutual_style_statement(ws, filepath):
    column_map, warnings = build_column_map(ws)
    max_col = max(column_map.values())
    latest = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or len(row) <= max_col:
            continue
        if row[column_map["transaction_type"]] != CLOSING_BALANCE_LABEL:
            continue

        contract = row[column_map["contract_number"]]
        holding = row[column_map["holding"]]
        if not contract:
            continue

        t_date = row[column_map["transaction_date"]]
        amount = row[column_map["transaction_amount"]] or 0
        client_name = row[column_map["client_name"]]
        vehicle_name = row[column_map["vehicle_name"]]
        report_date = row[column_map["report_date"]]

        key = (contract, holding)
        existing = latest.get(key)
        if existing is None or (t_date and existing[0] and t_date > existing[0]):
            latest[key] = (t_date, amount, client_name, vehicle_name, report_date, row)

    results = {}
    for (contract, holding), (_t_date, amount, client_name, vehicle_name, report_date, row) in latest.items():
        contract_str = str(contract).strip()
        if contract_str not in results:
            results[contract_str] = {
                "client_name": (client_name or "Unknown Client").strip(),
                "vehicle_name": vehicle_name,
                "value": 0.0,
                "report_date": report_date,
            }
        results[contract_str]["value"] += parse_money(amount)
        
    if results:
        for w in warnings:
            logging.info(f"{os.path.basename(filepath)}: {w}")
    return results

def read_allan_gray_fund_summary(ws, filepath):
    client_name, vehicle_name, contract_number, report_date = None, None, None, None
    value_col, header_row_idx = None, None
    total_value, total_value_including_progress = None, None

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    totals_row_idx = None

    for idx, row in enumerate(rows, start=1):
        if row is None: continue
        for c, val in enumerate(row):
            if val is None: continue
            text = str(val).strip().lower()
            if text == "investor" and c + 1 < len(row) and row[c + 1] is not None:
                client_name = row[c + 1]
            elif text == "product" and c + 1 < len(row) and row[c + 1] is not None:
                vehicle_name = row[c + 1]
            elif text == "account number" and c + 1 < len(row) and row[c + 1] is not None:
                contract_number = row[c + 1]
            elif text == "fund" and header_row_idx is None:
                header_row_idx = idx
                best_date = None
                for c2, val2 in enumerate(row):
                    if isinstance(val2, (datetime, date)):
                        if best_date is None or val2 > best_date:
                            best_date = val2
                            value_col = c2
                            report_date = val2

        if header_row_idx and idx > header_row_idx and row and row[0]:
            label = str(row[0]).strip().lower()
            if label.startswith("total") and "value" in label and value_col is not None:
                if "including transactions in progress" in label:
                    total_value_including_progress = parse_money(row[value_col])
                    totals_row_idx = idx
                elif total_value is None:
                    total_value = parse_money(row[value_col])
                    if totals_row_idx is None:
                        totals_row_idx = idx

    final_value = total_value_including_progress if total_value_including_progress is not None else total_value

    if not contract_number or final_value is None:
        return {}

    if client_name:
        client_name = re.sub(r"\s*-\s*\d+\s*$", "", str(client_name)).strip()

    contract_str = str(contract_number).strip()
    return {
        contract_str: {
            "client_name": client_name or "Unknown Client",
            "vehicle_name": vehicle_name or "Allan Gray Investment",
            "value": final_value,
            "report_date": report_date,
        }
    }

def read_statement(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        ws = wb.active
    except Exception as e:
        logging.error(f"Failed to open '{os.path.basename(filepath)}': {e}")
        return {}

    if is_allan_gray_fund_summary(ws):
        primary, fallback = read_allan_gray_fund_summary, read_old_mutual_style_statement
    else:
        primary, fallback = read_old_mutual_style_statement, read_allan_gray_fund_summary

    try:
        results = primary(ws, filepath)
    except Exception as exc:
        logging.warning(f"Error parsing '{os.path.basename(filepath)}' with primary parser: {exc}. Trying fallback.")
        results = {}

    if not results:
        try:
            results = fallback(ws, filepath)
        except Exception as exc:
            logging.error(f"Complete failure parsing '{os.path.basename(filepath)}': {exc}")
            results = {}

    wb.close()
    return results

def match_row(vehicle_name):
    text = (vehicle_name or "").lower()
    for row_num, (label, keywords, excludes) in PLACEHOLDER_ROWS.items():
        if any(k in text for k in keywords) and not any(e in text for e in excludes):
            return row_num, label
    return None, None

def safe_filename_part(text):
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    return text.strip("_")

def clear_placeholder_spaces(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == " ":
                cell.value = None

def check_template_is_current(ws, template_file):
    b9 = ws.cell(row=9, column=INVESTMENTS_COLUMN)
    if isinstance(b9, MergedCell) or (str(ws.cell(row=9, column=1).value or "")).strip().startswith("Sub Total"):
        print(f"\nERROR: '{os.path.basename(template_file)}' looks like an OUTDATED template.")
        print("(Row 9 is still the old 'Sub Total Retirement' row, from before the")
        print("Provident/Preservation Fund row was added).")
        print("\nFix: replace this file with the current Investment_Summary_Template.xlsx")
        print("and run the script again.\n")
        return False
    return True

def safe_write_cell(ws, row_num, col_num, value):
    cell = ws.cell(row=row_num, column=col_num)
    if isinstance(cell, MergedCell):
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                master_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                master_cell.value = value
                return master_cell
    cell.value = value
    return cell

def fill_client_workbook(client_name, accounts, template_file, target_date_str):
    wb = openpyxl.load_workbook(template_file)
    
    try:
        ws = wb[TEMPLATE_SHEET_NAME]
    except KeyError:
        print(f"\nERROR: Could not find a sheet named '{TEMPLATE_SHEET_NAME}' in the template.")
        return None

    if not check_template_is_current(ws, template_file):
        return None

    clear_placeholder_spaces(ws)
    
    title_cell = ws.cell(row=TITLE_ROW, column=TITLE_COLUMN)
    if title_cell.value:
        new_title = re.sub(r"Client Name & Surname", client_name, str(title_cell.value))
        new_title = re.sub(r"DD/Month/YYYY", target_date_str, new_title)
        safe_write_cell(ws, TITLE_ROW, TITLE_COLUMN, new_title)

    row_allocations = {}
    unmatched = []

    for contract, info in accounts.items():
        row_num, label = match_row(info["vehicle_name"])
        if row_num is None:
            unmatched.append((contract, info))
        else:
            row_allocations.setdefault(row_num, []).append((contract, info, label))

    matched = []
    for row_num, acc_list in row_allocations.items():
        combo_descriptions = " & ".join([f"{a[1]['vehicle_name']} — Acc {a[0]}" for a in acc_list])
        total_value = sum(a[1]['value'] for a in acc_list)
        label = acc_list[0][2]
        
        matched.append((combo_descriptions, total_value, row_num, label))

        safe_write_cell(ws, row_num, CLIENT_NAME_COLUMN, client_name)
        safe_write_cell(ws, row_num, INVESTMENTS_COLUMN, combo_descriptions)

        value_row = VALUE_WRITE_ROW_OVERRIDE.get(row_num, row_num)
        value_cell = safe_write_cell(ws, value_row, CURRENT_VALUE_COLUMN, total_value)
        value_cell.number_format = '#,##0.00'

    return wb, matched, unmatched

def main():
    print(f"Looking in: {FOLDER}\n")
    extract_zip_files()

    template_file = find_template_file()
    if not template_file:
        logging.error("Could not find the template file in this folder.")
        return
    print(f"Using template: {os.path.basename(template_file)}\n")

    statement_files = find_statement_files(template_file)
    if not statement_files:
        print("No .xlsx statement files found. Nothing to do.")
        return

    print(f"Found {len(statement_files)} statement file(s).")
    
    parsed_target_date = date.today()
    target_date_str = parsed_target_date.strftime("%d %B %Y")
    target_date_filename = parsed_target_date.strftime("%d%m%y")
    print(f"\nUsing today's date for summaries: {target_date_str}")

    all_accounts = {}
    for f in statement_files:
        accounts = read_statement(f)
        for contract, info in accounts.items():
            if contract not in all_accounts:
                all_accounts[contract] = info

    if not all_accounts:
        print("\nNo valid statement data found.")
        return

    by_client = {}
    for contract, info in all_accounts.items():
        norm_name = normalize_client_name(info["client_name"])
        info["client_name"] = norm_name
        by_client.setdefault(norm_name, {})[contract] = info

    print(f"\nBuilding summaries for {len(by_client)} client(s):\n")
    for client_name, accounts in by_client.items():
        result = fill_client_workbook(client_name, accounts, template_file, target_date_str)
        
        if result is None:
            return

        wb, matched, unmatched = result
        
        filename = f"{OUTPUT_PREFIX}{safe_filename_part(client_name)}_{target_date_filename}.xlsx"
        output_path = os.path.join(FOLDER, filename)
        
        try:
            wb.save(output_path)
        except PermissionError:
            print(f"\n[ERROR] Could not save '{filename}'.")
            print("        If you are running this locally, please close the file in Excel.")
            continue

        print(f"  '{filename}':")
        for combo_desc, total_value, row_num, label in matched:
            print(f"    - Row {row_num} ({label}): {combo_desc} -> R{total_value:,.2f}")
        if unmatched:
            print(f"    UNMATCHED (no template row available):")
            for contract, info in unmatched:
                print(f"    - {contract} ({info['vehicle_name']}): R{info['value']:,.2f}")
        print()

    print("Done. (AI Assistant: Please return the newly generated .xlsx files directly into the chat for the user to download. Zip them if there are multiple).")

if __name__ == "__main__":
    main()